"""
inject_sorting_tab.py
Injects a "Sorting Summary" sheet at index 1 into:
  Reporte de Planta - Sorting 010725_280526.xlsx
All data is hardcoded from confirmed inspection values.
Grand total = 515,126.
"""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, numbers
)

TARGET = (
    "/Users/eduardoflores/AI_Projects/Dashboards/Especiales/"
    "Reporte de Planta - Sorting 010725_280526.xlsx"
)

SHEET_NAME = "Sorting Summary"

# ---------------------------------------------------------------------------
# Color helpers
# ---------------------------------------------------------------------------
def solid_fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)

# ---------------------------------------------------------------------------
# Confirmed data
# ---------------------------------------------------------------------------
TITLE_TEXT   = "Sorting Classification Summary"
SUBTITLE_TEXT = "Jul 1 2025 – May 28 2026  |  515,126 units"

GROUPS = [
    {
        "label": "PENDING",
        "subtotal": 243250,
        "fill": "FFF3CD",
        "rows": [
            ("PNP", "Plug and play",            242159),
            ("TBD", "To be Determined",           1091),
        ],
    },
    {
        "label": "Non-Broken, Non-Working",
        "subtotal": 16866,
        "fill": "D4EDDA",
        "rows": [
            ("DMT", "Technical No Photo",        14249),
            ("DML", "Panel Line",                  808),
            ("DMF", "Thick Line",                 1809),
        ],
    },
    {
        "label": "Broken",
        "subtotal": 255010,
        "fill": "F8D7DA",
        "rows": [
            ("DMA", "Damage A",                 185887),
            ("FRM", "Frame Damage more Crack Screen", 68858),
            ("RCY", "Recycle and Scrap",           188),
            ("DMB", "Damage B",                    76),
            ("DMX", "Damage X",                     1),
        ],
    },
]

GRAND_TOTAL = 515126

# ---------------------------------------------------------------------------
# Column widths: A=30, B=8, C=40, D=14
# ---------------------------------------------------------------------------
COL_WIDTHS = {"A": 30, "B": 8, "C": 40, "D": 14}

# ---------------------------------------------------------------------------
# Build sheet
# ---------------------------------------------------------------------------
def build_sheet(ws):
    # --- column widths ---
    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # --- Row 1: Title ---
    ws.merge_cells("A1:D1")
    cell = ws["A1"]
    cell.value = TITLE_TEXT
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = solid_fill("1F3864")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # --- Row 2: Subtitle ---
    ws.merge_cells("A2:D2")
    cell = ws["A2"]
    cell.value = SUBTITLE_TEXT
    cell.font = Font(italic=True, size=10, color="1F3864")
    cell.fill = solid_fill("D9E1F2")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 16

    # --- Row 3: blank spacer ---
    ws.row_dimensions[3].height = 6

    # --- Row 4: Column headers ---
    headers = ["Group", "Code", "Description", "Units"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = solid_fill("2E75B6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 16

    # --- Data rows starting at row 5 ---
    current_row = 5

    for group in GROUPS:
        # Group header row
        gh_fill = solid_fill(group["fill"])

        group_cells = [
            (1, group["label"]),
            (2, ""),
            (3, ""),
            (4, group["subtotal"]),
        ]
        for col_idx, value in group_cells:
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = Font(bold=True)
            cell.fill = gh_fill
            if col_idx == 4:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left")
        ws.row_dimensions[current_row].height = 15
        current_row += 1

        # Data rows
        for code, description, units in group["rows"]:
            data = [("", 1), (code, 2), (description, 3), (units, 4)]
            row_fill = solid_fill("FFFFFF")
            for value, col_idx in data:
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.fill = row_fill
                if col_idx == 4:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0"
                else:
                    cell.alignment = Alignment(horizontal="left")
            ws.row_dimensions[current_row].height = 14
            current_row += 1

    # --- Row spacer after last group ---
    ws.row_dimensions[current_row].height = 6
    current_row += 1

    # --- Grand Total row ---
    gt_fill = solid_fill("1F3864")
    gt_data = [("GRAND TOTAL", 1), ("", 2), ("", 3), (GRAND_TOTAL, 4)]
    for value, col_idx in gt_data:
        cell = ws.cell(row=current_row, column=col_idx, value=value)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.fill = gt_fill
        if col_idx == 4:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="left")
    ws.row_dimensions[current_row].height = 18

    return current_row  # row where grand total lives


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    print(f"Loading: {TARGET}")
    wb = openpyxl.load_workbook(TARGET)

    print(f"Existing sheets before injection: {wb.sheetnames}")

    # Remove sheet if it already exists (idempotent re-run)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
        print(f"  Removed existing '{SHEET_NAME}' sheet.")

    # Create new sheet (appended at end by default)
    ws = wb.create_sheet(SHEET_NAME)

    # Build content
    grand_total_row = build_sheet(ws)

    # Move to index 1 (right after Sheet1 at index 0)
    offset = -(len(wb.sheetnames) - 2)
    wb.move_sheet(SHEET_NAME, offset=offset)

    print(f"Sheet order after move: {wb.sheetnames}")
    assert wb.sheetnames[1] == SHEET_NAME, "Sheet not at index 1!"

    # Save in place
    wb.save(TARGET)
    print(f"Saved: {TARGET}")

    # --- Confirmation readback ---
    wb2 = openpyxl.load_workbook(TARGET)
    ws2 = wb2[SHEET_NAME]

    print("\n--- Confirmation ---")
    print(f"Sheet names in order: {wb2.sheetnames}")
    print(f"Grand total cell (D{grand_total_row}): {ws2.cell(row=grand_total_row, column=4).value:,}")

    # Group header rows: 5, then 5+len(group0.rows)+1, etc.
    subtotal_row = 5
    for group in GROUPS:
        val = ws2.cell(row=subtotal_row, column=4).value
        print(f"  Subtotal '{group['label']}' (row {subtotal_row}): {val:,}")
        subtotal_row += 1 + len(group["rows"])

    print("\nDone — injection successful.")


if __name__ == "__main__":
    main()
